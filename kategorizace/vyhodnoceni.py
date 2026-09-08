"""
Vyhodnoceni presnosti klasifikatoru proti "gold" seznamu (rucni nebo LLM
kategorizace tychze firem).

    python3 -m kategorizace.vyhodnoceni vystup_kategorizace.csv gold.xlsx

gold muze byt primo drivejsi vystup dodavatele.py (ma sloupec "Kód kategorie"
vyplneny z LLM kroku). Parovani firem je podle nazvu (bez pravnich forem).
"""

import argparse
import csv
import os
import re
from collections import Counter

from .normalizace import bez_diakritiky, klic_nazvu
from .taxonomie import KATEGORIE, ict_relevance

_KOD_RE = re.compile(r"\b[A-Z]{2,4}-\d{2}\b")

_NAZEV_HLAVICKY = {"nazev", "název", "name", "jmeno", "jméno", "firma", "dodavatel",
                   "supplier", "company", "puvodni nazev", "původní název"}
_KOD_HLAVICKY = {"kod kategorie", "kód kategorie", "kategorie kod", "kategorie kód",
                 "kod", "kód", "code", "category code"}


def _kanon(h):
    h = bez_diakritiky((h or "").strip().lower())
    return re.sub(r"\s+", " ", re.sub(r"[^0-9a-z ]+", " ", h)).strip()


def _nacti_tabulku(cesta):
    pripona = os.path.splitext(cesta)[1].lower()
    if pripona in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("Cteni XLSX vyzaduje openpyxl - pouzijte .venv/bin/python.")
        wb = load_workbook(cesta, read_only=True, data_only=True)
        ws = wb.active
        radky = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return radky
    with open(cesta, encoding="utf-8-sig", newline="") as f:
        vzorek = f.read(8192)
        f.seek(0)
        try:
            d = csv.Sniffer().sniff(vzorek, delimiters=";,\t|")
        except csv.Error:
            d = csv.excel
            d.delimiter = ";" if ";" in vzorek else ","
        return [list(r) for r in csv.reader(f, dialect=d)]


def _sloupec(hlavicka, varianty):
    kanon = [_kanon(h) for h in hlavicka]
    chtene = {_kanon(v) for v in varianty}
    for i, k in enumerate(kanon):
        if k in chtene:
            return i
    return None


def _kod(text):
    m = _KOD_RE.search((text or "").upper())
    return m.group(0) if m else ""


def nacti_gold(cesta):
    radky = _nacti_tabulku(cesta)
    if not radky:
        return {}
    i_nazev = _sloupec(radky[0], _NAZEV_HLAVICKY) or 0
    i_kod = _sloupec(radky[0], _KOD_HLAVICKY)
    if i_kod is None:
        raise SystemExit("V gold souboru nenalezen sloupec s kodem kategorie "
                         "(napr. 'Kód kategorie').")
    gold = {}
    for r in radky[1:]:
        if i_nazev >= len(r):
            continue
        klic = klic_nazvu(r[i_nazev])
        kod = _kod(r[i_kod]) if i_kod < len(r) else ""
        if klic and kod:
            gold[klic] = kod
    return gold


def nacti_predikce(cesta):
    radky = _nacti_tabulku(cesta)
    if not radky:
        return {}
    h = radky[0]
    i_nazev = _sloupec(h, _NAZEV_HLAVICKY) or 0
    i_kod = _sloupec(h, _KOD_HLAVICKY)
    kanon = [_kanon(x) for x in h]
    i_rozh = kanon.index("rozhodnuti") if "rozhodnuti" in kanon else None
    i_alt = kanon.index("alternativy") if "alternativy" in kanon else None
    pred = {}
    for r in radky[1:]:
        if i_nazev >= len(r):
            continue
        klic = klic_nazvu(r[i_nazev])
        if not klic:
            continue
        kod = _kod(r[i_kod]) if i_kod is not None and i_kod < len(r) else ""
        rozh = r[i_rozh].strip().upper() if i_rozh is not None and i_rozh < len(r) else ""
        alt = []
        if i_alt is not None and i_alt < len(r):
            alt = _KOD_RE.findall((r[i_alt] or "").upper())
        top1 = alt[0] if alt else kod
        pred[klic] = {"kod": kod, "rozhodnuti": rozh, "top1": top1, "alt": alt}
    return pred


def _pct(a, b):
    return (100.0 * a / b) if b else 0.0


def _radek(popis, spravne, celkem):
    return "  %-42s %4d / %-4d  %5.1f %%" % (popis, spravne, celkem, _pct(spravne, celkem))


def vyhodnot(pred, gold):
    spolecne = sorted(set(pred) & set(gold))
    print("Firem v predikci: %d | v gold: %d | spolecnych: %d\n" % (len(pred), len(gold), len(spolecne)))
    if not spolecne:
        print("Zadna spolecna firma - zkontrolujte, ze jde o tyz seznam.")
        return

    celkem = len(spolecne)
    leaf_ok = grp_ok = top3_ok = 0
    auto = auto_ok = auto_grp_ok = 0
    over = over_ok = 0
    llm_by_slo = 0  # rozhodnuti LLM, ale top1 by trefil gold
    konfuze = Counter()
    skup_celkem = Counter()
    skup_ok = Counter()

    for klic in spolecne:
        g = gold[klic]
        p = pred[klic]
        g_grp = KATEGORIE.get(g, ["", ""])[0]
        t1 = p["top1"]
        t1_grp = KATEGORIE.get(t1, ["", ""])[0]

        skup_celkem[g_grp] += 1
        if t1 == g:
            leaf_ok += 1
            skup_ok[g_grp] += 1
        elif t1_grp and t1_grp == g_grp:
            skup_ok[g_grp] += 1
        if t1_grp == g_grp and t1_grp:
            grp_ok += 1
        if g in p["alt"][:3]:
            top3_ok += 1
        if t1 != g:
            konfuze[(g, t1)] += 1

        if p["rozhodnuti"] == "AUTO":
            auto += 1
            if p["kod"] == g:
                auto_ok += 1
            if KATEGORIE.get(p["kod"], ["", ""])[0] == g_grp:
                auto_grp_ok += 1
        elif p["rozhodnuti"] == "OVERIT":
            over += 1
            if p["kod"] == g:
                over_ok += 1
        elif p["rozhodnuti"] == "LLM" and t1 == g:
            llm_by_slo += 1

    print("PRESNOST TOP-1 (potencial klasifikatoru, bez ohledu na rozhodnuti)")
    print(_radek("list (kod kategorie)", leaf_ok, celkem))
    print(_radek("skupina", grp_ok, celkem))
    print(_radek("gold mezi top-3 navrhy", top3_ok, celkem))
    print()
    print("PODLE ROZHODNUTI (co by se stalo, kdyz AUTO prijmeme napevno)")
    print(_radek("AUTO - pokryti ze vsech firem", auto, celkem))
    print(_radek("AUTO - list spravne", auto_ok, auto))
    print(_radek("AUTO - aspon spravna skupina", auto_grp_ok, auto))
    print(_radek("OVERIT - list spravne", over_ok, over))
    print("  LLM rozhodnuti, ktere by top-1 presto trefil gold: %d" % llm_by_slo)
    print()
    print("PRESNOST PODLE SKUPINY (list nebo spravna skupina / celkem)")
    for grp in sorted(skup_celkem, key=lambda x: -skup_celkem[x]):
        print(_radek(grp, skup_ok[grp], skup_celkem[grp]))
    print()
    # ICT relevance (vstup pro navazujici krok ISO 27001)
    ict_c = ict_ok = ict_auto = ict_auto_ok = 0
    for klic in spolecne:
        gi = ict_relevance(gold[klic])
        pi = ict_relevance(pred[klic]["top1"])
        if not gi:
            continue
        ict_c += 1
        if gi == pi:
            ict_ok += 1
        if pred[klic]["rozhodnuti"] == "AUTO":
            ict_auto += 1
            if ict_relevance(pred[klic]["kod"]) == gi:
                ict_auto_ok += 1
    if ict_c:
        print("ICT RELEVANCE (ano / hranicni / ne - pro navazujici krok ISO 27001)")
        print(_radek("shoda ICT priznaku (vsechny firmy)", ict_ok, ict_c))
        print(_radek("shoda ICT priznaku (jen AUTO)", ict_auto_ok, ict_auto))
        print()

    print("NEJCASTEJSI ZAMENY  (gold -> navrzeno)")
    for (g, t), n in konfuze.most_common(15):
        gn = KATEGORIE.get(g, ["", g])[1]
        tn = KATEGORIE.get(t, ["", t])[1] if t else "(zadny navrh)"
        print("  %2d x  %s %-38s -> %s %s" % (n, g, gn[:38], t or "----", tn))


def main(argv=None):
    p = argparse.ArgumentParser(
        prog="kategorizace.vyhodnoceni",
        description="Porovna vystup klasifikatoru s gold seznamem (rucni/LLM kategorizace).",
    )
    p.add_argument("predikce", help="CSV vystup z kategorizuj.py")
    p.add_argument("gold", help="CSV/XLSX se sloupci Nazev a Kod kategorie (napr. drivejsi vystup)")
    a = p.parse_args(argv)
    vyhodnot(nacti_predikce(a.predikce), nacti_gold(a.gold))


if __name__ == "__main__":
    main()
