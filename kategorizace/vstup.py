"""
Nezavisle cteni vstupniho seznamu firem (CSV / XLSX / TXT). Zamerne nesdili
kod s dodavatele.py - modul ma bezet samostatne.

Vraci seznam dict: {"nazev","mesto","zeme","nace","ico","web"} (chybejici = "").
"""

import csv
import os
import re

from .normalizace import bez_diakritiky

_SLOUPCE = {
    "nazev": {"nazev", "název", "name", "jmeno", "jméno", "firma", "spolecnost",
              "společnost", "dodavatel", "supplier", "company", "obchodni jmeno",
              "obchodní jméno"},
    "mesto": {"mesto", "město", "city", "obec", "sidlo", "sídlo"},
    "zeme": {"zeme", "země", "country", "stat", "stát", "kod zeme", "kód země"},
    "nace": {"nace", "nace kod", "nace kód", "obor", "cinnost", "činnost"},
    "ico": {"ico", "ičo", "ic", "identifikacni cislo", "identifikační číslo"},
    "web": {"web", "www", "url", "webstranka", "webová stránka", "webova stranka",
            "stranky", "stránky", "website", "domena", "doména", "domain"},
}


def _kanon(hlavicka):
    h = bez_diakritiky((hlavicka or "").strip().lower())
    h = re.sub(r"[^0-9a-z ]+", " ", h)
    return re.sub(r"\s+", " ", h).strip()


def _mapuj_hlavicku(hlavicky):
    mapa = {}
    for i, h in enumerate(hlavicky):
        k = _kanon(h)
        for cil, varianty in _SLOUPCE.items():
            if cil in mapa:
                continue
            if k in {_kanon(v) for v in varianty}:
                mapa[cil] = i
    return mapa


def _radky_z_tabulky(radky):
    if not radky:
        return []
    hlavicka = radky[0]
    mapa = _mapuj_hlavicku(hlavicka)
    if "nazev" not in mapa:
        # bez rozpoznane hlavicky bereme prvni sloupec jako nazev
        mapa = {"nazev": 0}
        datove = radky
    else:
        datove = radky[1:]
    out = []
    for r in datove:
        def bunka(klic):
            i = mapa.get(klic)
            return (r[i].strip() if i is not None and i < len(r) and r[i] else "")
        nazev = bunka("nazev")
        if not nazev:
            continue
        out.append({
            "nazev": nazev,
            "mesto": bunka("mesto"),
            "zeme": bunka("zeme"),
            "nace": bunka("nace"),
            "ico": bunka("ico"),
            "web": bunka("web"),
        })
    return out


def _cti_csv(cesta):
    with open(cesta, encoding="utf-8-sig", newline="") as f:
        vzorek = f.read(8192)
        f.seek(0)
        try:
            dialekt = csv.Sniffer().sniff(vzorek, delimiters=";,\t|")
        except csv.Error:
            dialekt = csv.excel
            dialekt.delimiter = ";" if ";" in vzorek else ","
        return [list(r) for r in csv.reader(f, dialect=dialekt)]


def _cti_xlsx(cesta):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit(
            "Cteni XLSX vyzaduje openpyxl (pip install openpyxl) nebo pouzijte "
            ".venv/bin/python. CSV funguje bez nej."
        )
    wb = load_workbook(cesta, read_only=True, data_only=True)
    ws = wb.active
    radky = []
    for r in ws.iter_rows(values_only=True):
        radky.append(["" if b is None else str(b) for b in r])
    wb.close()
    return radky


def _cti_txt(cesta):
    with open(cesta, encoding="utf-8-sig") as f:
        return [[radek.strip()] for radek in f if radek.strip()]


def nacti(cesta):
    pripona = os.path.splitext(cesta)[1].lower()
    if pripona in (".xlsx", ".xlsm"):
        radky = _cti_xlsx(cesta)
    elif pripona == ".txt":
        radky = _cti_txt(cesta)
    else:
        radky = _cti_csv(cesta)
    return _radky_z_tabulky(radky)
