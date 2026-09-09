#!/usr/bin/env python3
"""
Nahled sesitu bez site - ukazkova data misto realneho dohledani.

Slouzi k posouzeni podoby vystupu (rozvrzeni listu, karty, barvy, vzorce),
ne k praci s realnymi dodavateli. Vzniknou dva soubory:

    test_vystup/dodavatele.xlsx            jak sesit prijde ke klientovi
    test_vystup/dodavatele-vyplneny.xlsx   tyz sesit s ukazkove vyplnenymi
                                           rucnimi sloupci, aby bylo videt,
                                           jak vypada spoctena kriticnost
                                           a zvyrazneni chybejicich udaju

    python3 nahled.py [adresar]
"""

import os
import sys
from datetime import date, timedelta

import vystup
import zdroje
from zdroje import (IDENT_NAZEV, IDENT_NEOVERENO, IDENT_REJSTRIK, IDENT_VIES,
                    Firma)

ADRESAR = "test_vystup"

# (kod kreditora, nazev, identita, ico, dic, zeme, mesto, web, nace,
#  kod kategorie, co dodava, jistota, obory, poznamky)
UKAZKA = [
    ("D0001", "Alza.cz a.s.", IDENT_REJSTRIK, "27082440", "CZ27082440", "CZ",
     "Praha", "https://www.alza.cz", ["47910"],
     "ICT-01", "e-shop s výpočetní technikou a elektronikou", "vysoká",
     ["Výroba, obchod a služby neuvedené v přílohách 1 až 3 živnostenského zákona"],
     []),
    ("D0002", "GABEN, spol. s r. o.", IDENT_REJSTRIK, "19012021", "CZ19012021", "CZ",
     "Ostrava", "https://www.gaben.cz", ["46520", "62090"],
     "ICT-08", "čárové kódy, RFID a systémy pro sběr dat", "vysoká",
     ["Poskytování software, poradenství v oblasti informačních technologií"],
     []),
    ("D0003", "WAGO Kontakttechnik GmbH & Co. KG", IDENT_VIES, "", "DE811150595",
     "DE", "Minden", "https://www.wago.com", [],
     "ELE-04", "svorkovnice, konektory a automatizační technika", "vysoká",
     [], []),
    ("D0004", "ESET, spol. s r.o.", IDENT_REJSTRIK, "31333532", "SK2020317068",
     "SK", "Bratislava", "https://www.eset.com", ["62010"],
     "ICT-05", "antivirové a bezpečnostní řešení pro koncové stanice", "vysoká",
     [], []),
    ("D0005", "Cloudflare, Inc.", IDENT_NEOVERENO, "", "", "US",
     "San Francisco", "https://www.cloudflare.com", [],
     "ICT-04", "CDN, ochrana proti DDoS a DNS služby", "vysoká",
     [], ["země bez veřejného rejstříku zdarma – identita neověřena"]),
    ("D0006", "Vodafone Czech Republic a.s.", IDENT_REJSTRIK, "25788001",
     "CZ25788001", "CZ", "Praha", "https://www.vodafone.cz", ["61200"],
     "ICT-06", "mobilní hlasové a datové služby, firemní tarify",
     "vysoká", ["Výroba, obchod a služby neuvedené v přílohách 1 až 3"], []),
    ("D0007", "ČEZ Prodej, a.s.", IDENT_REJSTRIK, "27232433", "CZ27232433", "CZ",
     "Praha", "https://www.cez.cz", ["35140"],
     "FAC-05", "dodávka elektřiny a plynu do provozovny", "vysoká", [], []),
    ("D0008", "SECURITAS ČR s.r.o.", IDENT_REJSTRIK, "43872026", "CZ43872026",
     "CZ", "Praha", "https://www.securitas.cz", ["80100"],
     "FAC-09", "fyzická ostraha areálu a recepční služby", "vysoká", [], []),
    ("D0009", "Úklidový servis Novák s.r.o.", IDENT_NAZEV, "", "", "CZ",
     "Brno", "", [],
     "PROV-05", "pravidelný úklid kanceláří a výrobní haly", "střední", [],
     ["web nenalezen", "identita jen podle shody názvu – ověřte IČO"]),
    ("D0010", "HAVEL & PARTNERS s.r.o.", IDENT_REJSTRIK, "26454807", "CZ26454807",
     "CZ", "Praha", "https://www.havelpartners.cz", ["69100"],
     "PRO-02", "právní služby, smluvní agenda a compliance", "vysoká", [], []),
    ("D0011", "ABRA Software a.s.", IDENT_REJSTRIK, "25097563", "CZ25097563",
     "CZ", "Praha", "https://www.abra.eu", ["62010"],
     "ICT-02", "ERP systém, licence a implementace", "vysoká", [], []),
    ("D0012", "TNT Express Worldwide, spol. s r.o.", IDENT_REJSTRIK, "15888484",
     "CZ15888484", "CZ", "Praha", "https://www.tnt.com", ["53200"],
     "LOG-02", "expresní kurýrní přeprava zásilek", "vysoká", [], []),
    ("D0013", "Papírnictví Kolman", IDENT_NEOVERENO, "", "", "CZ",
     "", "", [],
     "PROV-03", "kancelářské potřeby", "nízká", [],
     ["v rejstříku nenalezeno – nejspíš OSVČ nebo zkrácený název"]),
    ("D0014", "Testovací dodavatel bez zařazení", IDENT_NEOVERENO, "", "", "",
     "", "", [],
     "", "", "", [], ["ke kategorizaci nedorazila odpověď z chatu"]),
    ("D0015", "Jan Novák", IDENT_REJSTRIK, "00719331", "", "CZ",
     "Ostrožská Lhota", "", [],
     "", "", "", [], []),
]

# Kdo je fyzicka osoba (pravni forma z ARES). U D0015 na tom stoji ukazka
# znacky XXX-01: OSVC neni v obchodnim rejstriku, takze zadny zapsany
# predmet podnikani neexistuje a jmeno cloveka obor neprozradi.
FORMY = {"D0015": "101"}

# Indexy do vystup.PRISTUP a vystup.NAHRADITELNOST - psat je cislem by
# znamenalo, ze se ukazka tise rozejde s poradim v nabidce.
(BEZ_PRISTUPU, FYZICKY, PREDANI, ZPRACOVANI,
 UZIVATEL, SPRAVA, OBOJI) = range(len(vystup.PRISTUP))
BEZNE, OBTIZNE, ZAVISLOST = range(len(vystup.NAHRADITELNOST))

# Ukazkove hodnoty rucnich sloupcu pro druhy soubor. Klic = kod kreditora,
# hodnota = (pristup, nahraditelnost, vlastnik, posouzeni, prezkoumani, pozn.)
VYPLNENO = {
    "D0001": (BEZ_PRISTUPU, BEZNE, "J. Dvořák (nákup)", -120, 610, ""),
    "D0002": (UZIVATEL, OBTIZNE, "M. Horák (výroba)", -60, 300,
              "Přístup do skladového systému přes API."),
    "D0003": (BEZ_PRISTUPU, OBTIZNE, "M. Horák (výroba)", -200, 530, ""),
    "D0004": (ZPRACOVANI, BEZNE, "P. Nový (IT)", -30, 335,
              "DPA podepsáno, telemetrie mimo EU."),
    "D0005": (ZPRACOVANI, ZAVISLOST, "P. Nový (IT)", -30, 335,
              "Provoz webu stojí a padá s nimi, migrace by trvala týdny."),
    "D0006": (ZPRACOVANI, OBTIZNE, "P. Nový (IT)", -400, -30,
              "Přezkoumání propadlo, řeší se."),
    "D0007": (BEZ_PRISTUPU, OBTIZNE, "", None, None, ""),
    "D0008": (FYZICKY, BEZNE, "L. Marek (provoz)", -90, 275, ""),
    "D0009": (FYZICKY, BEZNE, "", None, None,
              "Vstup do prostor mimo pracovní dobu."),
    "D0010": (PREDANI, OBTIZNE, "K. Beránková (právní)", -150, 580,
              "Dostává smluvní a personální dokumenty, do systémů nechodí."),
    "D0011": (SPRAVA, ZAVISLOST, "", None, None,
              "Vzdálená správa ERP včetně přístupu do databáze – chybí vlastník."),
    "D0012": (FYZICKY, BEZNE, "L. Marek (provoz)", -220, 500, ""),
    "D0013": (BEZ_PRISTUPU, BEZNE, "", None, None, ""),
    "D0014": (None, None, "", None, None, ""),
}


def _firmy():
    out = []
    for (kod, nazev, identita, ico, dic, zeme, mesto, web, nace,
         _k, _c, _j, obory, poznamky) in UKAZKA:
        f = Firma(vstup_kod=kod, vstup_nazev=nazev, vstup_ico=ico,
                  vstup_dic=dic, vstup_zeme=zeme)
        f.nazev, f.ico, f.dic = nazev, ico, dic
        f.zeme, f.mesto, f.web = zeme, mesto, web
        f.nace, f.obory, f.identita = list(nace), list(obory), identita
        f.pravni_forma = FORMY.get(kod, "112")
        f.poznamky = list(poznamky)
        if zdroje.bez_zarazeni(f):
            f.poznamky.append("fyzická osoba – v obchodním rejstříku není, "
                              "žádný podklad k oboru se nenašel")
        if web:
            f.web_text = ("Oficiální web dodavatele – ukázkový text, "
                          "v ostrém běhu je tu úvodní odstavec ze stránky.")
        out.append(f)
    return out


def _odpovedi():
    return {i: {"kod": r[9], "popis": r[10], "jistota": r[11]}
            for i, r in enumerate(UKAZKA, 1) if r[9]}


def _vypln(cesta, firmy):
    """Doplni do hotoveho sesitu ukazkove rucni hodnoty."""
    from openpyxl import load_workbook

    wb = load_workbook(cesta)
    ws = wb["Dodavatelé"]
    # zahlavi je pod kartou - najdeme ho podle prvniho sloupce
    zahlavi = next(r for r in range(1, 30) if ws.cell(r, 1).value == "Kreditor")
    dnes = date.today()
    for i, f in enumerate(firmy, 1):
        r = zahlavi + i
        pristup, nahrad, vlastnik, od, do, pozn = VYPLNENO.get(f.vstup_kod,
                                                               (None,) * 6)
        if pristup is not None:
            ws.cell(r, vystup.S_PRISTUP, vystup.PRISTUP[pristup])
        if nahrad is not None:
            ws.cell(r, vystup.S_NAHRAD, vystup.NAHRADITELNOST[nahrad])
        if vlastnik:
            ws.cell(r, vystup.S_VLASTNIK, vlastnik)
        for sl, posun in ((vystup.S_POSOUZENI, od), (vystup.S_PREZKOUM, do)):
            if posun is not None:
                ws.cell(r, sl, dnes + timedelta(days=posun))
                ws.cell(r, sl).number_format = "DD.MM.YYYY"
        if pozn:
            ws.cell(r, vystup.S_POZNAMKA, pozn)
    wb.save(cesta)


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    adresar = argv[0] if argv else ADRESAR
    os.makedirs(adresar, exist_ok=True)

    firmy = _firmy()
    prazdny = os.path.join(adresar, "dodavatele.xlsx")
    vyplneny = os.path.join(adresar, "dodavatele-vyplneny.xlsx")

    vystup.zapis_excel(firmy, _odpovedi(), prazdny)
    vystup.zapis_excel(firmy, _odpovedi(), vyplneny)
    _vypln(vyplneny, firmy)
    cesty = vystup.zapis_davky(firmy, os.path.join(adresar, "davky"))

    print("Náhled z ukázkových dat (%d dodavatelů):" % len(firmy))
    print("  %s   jak sešit přijde ke klientovi" % prazdny)
    print("  %s   s ukázkově vyplněnými ručními sloupci" % vyplneny)
    print("  %s   ukázka dávky pro chat" % cesty[0])
    return 0


if __name__ == "__main__":
    sys.exit(main())
