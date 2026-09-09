"""
Kontrola vzorcu v Excelu. Vzorce nejde spustit bez Excelu, tak se tu jejich
logika prepocita v Pythonu podle stejne tabulky pravidel a porovna s tim, co
ma vyjit. Chrani pred tichou zmenou pravidel kritickosti.

    python3 test_vzorce.py
"""

import re
import sys

import vystup
from taxonomie import KATEGORIE, ict_relevance


ZADNY, FYZICKY, PREDANI, ZPRACOVANI, UZIVATEL, SPRAVA, OBOJI = vystup.PRISTUP
BEZNE, OBTIZNE, ZAVISLOST = vystup.NAHRADITELNOST


def kriticnost(ict, pristup, nahraditelnost, kod="ICT-01"):
    """Tataz pravidla jako _vzorec_kriticnost - drzet synchronne."""
    # volby, u kterych se dodavatel dostane k informacim nebo do systemu
    K_INFORMACIM = (PREDANI, ZPRACOVANI, UZIVATEL, SPRAVA, OBOJI)
    if not kod:
        return "nezařazeno"
    if not pristup:
        return "⟵ doplňte přístup"
    if (pristup == SPRAVA
            or nahraditelnost == ZAVISLOST
            or (ict == "ano" and pristup in K_INFORMACIM)):
        return "KRITICKÝ"
    if ict == "ano" or pristup in K_INFORMACIM or nahraditelnost == OBTIZNE:
        return "VÝZNAMNÝ"
    return "BĚŽNÝ"


PRIPADY = [
    # (ICT relevance, přístup, nahraditelnost) -> očekáváno
    (("ano", SPRAVA, BEZNE), "KRITICKÝ"),
    (("ne", SPRAVA, BEZNE), "KRITICKÝ"),      # správa systémů rozhoduje i mimo ICT
    (("ano", ZPRACOVANI, BEZNE), "KRITICKÝ"),  # ICT dodávka + naše data u něj
    (("ne", ZPRACOVANI, BEZNE), "VÝZNAMNÝ"),   # cloud u neICT dodavatele ještě ne
    (("ano", UZIVATEL, BEZNE), "KRITICKÝ"),    # ICT dodávka + účet v našich systémech
    (("ne", UZIVATEL, BEZNE), "VÝZNAMNÝ"),
    (("ano", PREDANI, BEZNE), "KRITICKÝ"),     # ICT dodavatel, kterému posíláme data
    (("ne", PREDANI, BEZNE), "VÝZNAMNÝ"),      # samotné předání dat = významný
    (("ano", OBOJI, BEZNE), "KRITICKÝ"),
    (("ne", OBOJI, BEZNE), "VÝZNAMNÝ"),
    (("ne", ZADNY, ZAVISLOST), "KRITICKÝ"),   # single point of failure
    (("ano", ZADNY, BEZNE), "VÝZNAMNÝ"),
    (("ne", ZADNY, OBTIZNE), "VÝZNAMNÝ"),
    (("ne", FYZICKY, BEZNE), "BĚŽNÝ"),
    (("ano", FYZICKY, BEZNE), "VÝZNAMNÝ"),    # ICT dodávka bez přístupu k datům
    (("ne", FYZICKY, OBTIZNE), "VÝZNAMNÝ"),
    (("ne", ZADNY, BEZNE), "BĚŽNÝ"),
    (("ano", "", BEZNE), "⟵ doplňte přístup"),  # bez ručního vstupu nepočítá
    (("ne", ZADNY, BEZNE, ""), "nezařazeno"),   # bez kódu kategorie taky ne
]


def _zkontroluj_syntaxi(vzorec, jmeno):
    """Zavorky a uvozovky musi sedet - preklep by Excel ohlasil az u uzivatele."""
    chyby = []
    if vzorec.count("(") != vzorec.count(")"):
        chyby.append("%s: nesedí závorky" % jmeno)
    if vzorec.count('"') % 2:
        chyby.append("%s: lichý počet uvozovek" % jmeno)
    if not vzorec.startswith("="):
        chyby.append("%s: nezačíná '='" % jmeno)
    return chyby


def _test_nacitani_odpovedi():
    """Čtení odpovědí z chatu: formát, tolerance k šumu, kontrola párování."""
    import os
    import tempfile

    class _F:
        def __init__(self, n):
            self.vstup_nazev = self.nazev = n

    firmy = [_F("Alza.cz a.s."), _F("GABEN, spol. s r. o."), _F("ESET, spol. s r.o.")]
    chyby = []

    def _precti(obsah, s_kontrolou=True):
        with tempfile.TemporaryDirectory() as d:
            with open(os.path.join(d, "o.csv"), "w", encoding="utf-8") as f:
                f.write(obsah)
            return vystup.nacti_odpovedi(d, firmy if s_kontrolou else None)

    # správně spárovaná odpověď
    odp, nes = _precti("1;Alza.cz a.s.;OBH-01;e-shop;vysoká\n"
                       "2;GABEN, spol. s r. o.;ICT-08;RFID;vysoká\n")
    if len(odp) != 2 or nes:
        chyby.append("správná odpověď: %d řádků, %d nesouladů" % (len(odp), len(nes)))
    if odp.get(1, {}).get("kod") != "OBH-01" or odp.get(1, {}).get("popis") != "e-shop":
        chyby.append("špatně rozparsovaný řádek: %r" % odp.get(1))

    # posunuté číslování se musí zachytit, ne tiše přiřadit
    odp, nes = _precti("1;GABEN, spol. s r. o.;ICT-08;RFID;vysoká\n")
    if odp or len(nes) != 1:
        chyby.append("posunuté ID se nezachytilo (%d řádků, %d nesouladů)"
                     % (len(odp), len(nes)))

    # starší formát bez názvu musí pořád fungovat
    odp, nes = _precti("1;OBH-01;e-shop;vysoká\n")
    if odp.get(1, {}).get("kod") != "OBH-01":
        chyby.append("formát bez názvu firmy se nepřečetl: %r" % odp)

    # hlavička a povídání okolo se přeskočí
    odp, nes = _precti("Tady je výsledek:\nID;Název;Kód;Co dodává;Jistota\n"
                       "3;ESET, spol. s r.o.;ICT-05;antivirus;vysoká\n")
    if odp.get(3, {}).get("kod") != "ICT-05":
        chyby.append("šum okolo odpovědi se nepřeskočil: %r" % odp)

    # drobná odchylka v názvu se tolerovat má (chat občas zkrátí)
    odp, nes = _precti("1;Alza.cz;OBH-01;e-shop;vysoká\n")
    if not odp:
        chyby.append("zkrácený název se neuznal jako shoda")
    return chyby


def _test_osvc():
    """
    Fyzicka osoba, ke ktere se nenaslo nic: nesmi jit do davky pro chat
    (model by z holeho jmena vyrobil dohad) a v sesitu ma mit znacku XXX-01,
    ne prazdno ani kategorii od modelu.
    """
    import os
    import tempfile

    from zdroje import Firma, bez_zarazeni

    chyby = []
    osvc = Firma(vstup_kod="D001", vstup_nazev="Jan Novák", nazev="Jan Novák",
                 pravni_forma="101")
    osvc_s_nace = Firma(vstup_kod="D002", vstup_nazev="Petr Malý",
                        nazev="Petr Malý", pravni_forma="101", nace=["4779"])
    firma = Firma(vstup_kod="D003", vstup_nazev="ACME s.r.o.", nazev="ACME s.r.o.",
                  pravni_forma="112")
    firmy = [osvc, osvc_s_nace, firma]

    if not bez_zarazeni(osvc):
        chyby.append("OSVČ bez podkladů se nepoznala")
    for f in (osvc_s_nace, firma):
        if bez_zarazeni(f):
            chyby.append("%r se označila jako nezařaditelná" % f.vstup_nazev)

    with tempfile.TemporaryDirectory() as d:
        vystup.zapis_davky(firmy, d)
        text = "".join(open(os.path.join(d, j), encoding="utf-8").read()
                       for j in os.listdir(d))
    if "Jan Novák" in text:
        chyby.append("OSVČ bez podkladů se poslala do dávky pro chat")
    for jmeno in ("Petr Malý", "ACME s.r.o."):
        if jmeno not in text:
            chyby.append("%r z dávky vypadl, i když podklad má" % jmeno)
    # znacka nastroje nepatri do nabidky pro model - zacal by ji pouzivat sam
    if vystup.KOD_OSVC in text:
        chyby.append("%s se nabízí modelu v číselníku" % vystup.KOD_OSVC)
    # cislovani firem musi zustat na indexu, jinak se odpovedi sparuji spatne
    if "\n2 | Petr Malý" not in text or "\n3 | ACME s.r.o." not in text:
        chyby.append("vynechaná firma posunula číslování ostatních")

    try:
        from openpyxl import load_workbook
    except ImportError:
        return chyby
    with tempfile.TemporaryDirectory() as d:
        cesta = os.path.join(d, "t.xlsx")
        # model o firme 1 nic nevratil, u firmy 3 vratil kategorii
        vystup.zapis_excel(firmy, {3: {"kod": "PRO-01", "popis": "audit",
                                       "jistota": "vysoká"}}, cesta)
        ws = load_workbook(cesta)["Dodavatelé"]
        zahlavi = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Kreditor")
        kody = [ws.cell(zahlavi + i, vystup.S_KOD).value for i in (1, 2, 3)]
    if kody[0] != vystup.KOD_OSVC:
        chyby.append("OSVČ bez podkladů nedostala %s, ale %r"
                     % (vystup.KOD_OSVC, kody[0]))
    if kody[2] != "PRO-01":
        chyby.append("firmě se přepsal kód od modelu: %r" % kody[2])
    return chyby


def _test_sesit():
    """
    Vygeneruje sesit z par firem a overi, ze vzorce ukazuji na radek, na
    kterem opravdu jsou. Kvuli kartam nad tabulkou data nezacinaji na radku 2
    a posun by se jinak projevil az u klienta v Excelu.
    """
    import os
    import tempfile

    from zdroje import Firma

    chyby = []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ["openpyxl chybí – kontrola sešitu se nespustila"]

    # druha firma nema podklady, jen poznamku - prave u ni se musi ukazat,
    # ze se poznamka na Podklady dostane i bez jedineho nalezeneho zdroje
    firmy = [Firma(vstup_kod="D001", vstup_nazev="Alza.cz a.s."),
             Firma(vstup_nazev="Firma bez kódu kreditora",
                   poznamky=["web nenalezen"])]
    with tempfile.TemporaryDirectory() as d:
        cesta = os.path.join(d, "t.xlsx")
        vystup.zapis_excel(firmy, {1: {"kod": "ICT-01", "popis": "HW",
                                       "jistota": "vysoká"}}, cesta)
        wb = load_workbook(cesta)
        if wb.sheetnames[0] != "Úvod":
            chyby.append("první list je %r, čekán Úvod" % wb.sheetnames[0])
        for jmeno in ("Dodavatelé", "Číselník", "Souhrn", "Podklady", "Metodika"):
            if jmeno not in wb.sheetnames:
                chyby.append("v sešitu chybí list %s" % jmeno)
        ws = wb["Dodavatelé"]
        zahlavi = next((r for r in range(1, 40)
                        if ws.cell(r, 1).value == "Kreditor"), None)
        if zahlavi is None:
            return chyby + ["v listu Dodavatelé se nenašlo záhlaví"]
        if zahlavi < 2:
            chyby.append("nad tabulkou chybí karta (záhlaví je na řádku %d)" % zahlavi)
        for i in range(1, len(firmy) + 1):
            r = zahlavi + i
            vzorec = ws.cell(r, vystup.S_KRIT).value or ""
            if "%d" % r not in vzorec or not vzorec.startswith("="):
                chyby.append("kritičnost na řádku %d neodkazuje na svůj řádek: %r"
                             % (r, vzorec[:60]))
        if ws.cell(zahlavi + 1, vystup.S_KREDITOR).value != "D001":
            chyby.append("kód kreditora ze vstupu se nedostal do prvního sloupce")
        if ws.cell(zahlavi + 2, vystup.S_KREDITOR).value != 2:
            chyby.append("bez kódu kreditora se nedoplnilo pořadové číslo")
        # karta nesmi zasahovat do zahlavi ani prepsat data
        if ws.cell(zahlavi - 1, 1).value:
            chyby.append("mezi kartou a záhlavím chybí prázdný řádek")

        # poznamky nastroje patri k dohledavani, ne do evidence dodavatele -
        # maji byt na Podkladech, a to i u firmy, o ktere se nic nenaslo
        hlavicky = [vystup.SLOUPCE[i][0] for i in range(len(vystup.SLOUPCE))]
        for nechtene in ("Poznámky nástroje", "Objem ze vstupu"):
            if nechtene in hlavicky:
                chyby.append("sloupec %r má být z listu Dodavatelé pryč" % nechtene)
        ws_p = wb["Podklady"]
        zahlavi_p = next((r for r in range(1, 40)
                          if ws_p.cell(r, 1).value == "Kreditor"), None)
        if zahlavi_p is None:
            chyby.append("v listu Podklady se nenašlo záhlaví")
        else:
            posledni_sl = ws_p.max_column
            if ws_p.cell(zahlavi_p, posledni_sl).value != "Poznámky nástroje":
                chyby.append("Poznámky nástroje nejsou posledním sloupcem Podkladů")
            texty = [ws_p.cell(r, posledni_sl).value
                     for r in range(zahlavi_p + 1, ws_p.max_row + 1)]
            if not any(t and "web" in t.lower() for t in texty):
                chyby.append("poznámka nástroje se na Podklady nepřenesla: %r" % texty)

        # datum vygenerovani na kosilce byt nema - sesit se pouziva dlouhodobe
        ws_u = wb["Úvod"]
        for radek in ws_u.iter_rows(max_row=12, max_col=2):
            for bunka in radek:
                if isinstance(bunka.value, str) and "vygenerováno" in bunka.value:
                    chyby.append("na Úvodu zůstalo datum vygenerování")
        # text kosilky se slucuje pres A:B - co se tam nevejde, Excel oreze
        sirka = sum(ws_u.column_dimensions[sl].width for sl in ("A", "B"))
        for r in range(1, ws_u.max_row + 1):
            text = ws_u.cell(r, 1).value
            if (isinstance(text, str) and not ws_u.cell(r, 2).value
                    and len(text) > sirka):
                chyby.append("řádek %d Úvodu (%d znaků) se do %d nevejde: %r"
                             % (r, len(text), sirka, text[:40]))
    return chyby


def main():
    chyby = _test_nacitani_odpovedi() + _test_sesit() + _test_osvc()

    # 1) logika kritičnosti
    for vstup, cekano in PRIPADY:
        skutecnost = kriticnost(*vstup)
        if skutecnost != cekano:
            chyby.append("kritičnost%s = %r, čekáno %r" % (vstup, skutecnost, cekano))

    # 2) syntaxe vzorců
    for jmeno, vzorec in (("kritičnost", vystup._vzorec_kriticnost(2)),
                          ("ISO režim", vystup._vzorec_iso(2)),
                          ("VLOOKUP", vystup._vlookup(2, 2))):
        chyby += _zkontroluj_syntaxi(vzorec, jmeno)

    # 3) vzorce musí odkazovat na sloupce, které opravdu existují
    pocet = len(vystup.SLOUPCE)
    for jmeno, sl in (("KREDITOR", vystup.S_KREDITOR), ("KÓD", vystup.S_KOD),
                      ("KATEGORIE", vystup.S_KATEG), ("SKUPINA", vystup.S_SKUP),
                      ("ICT", vystup.S_ICT), ("POPIS", vystup.S_POPIS),
                      ("PŘÍSTUP", vystup.S_PRISTUP), ("NAHRADITELNOST", vystup.S_NAHRAD),
                      ("KRITIČNOST", vystup.S_KRIT), ("ISO", vystup.S_ISO),
                      ("VLASTNÍK", vystup.S_VLASTNIK),
                      ("POSOUZENÍ", vystup.S_POSOUZENI),
                      ("PŘEZKOUMÁNÍ", vystup.S_PREZKOUM),
                      ("POZNÁMKA", vystup.S_POZNAMKA)):
        if not 1 <= sl <= pocet:
            chyby.append("index sloupce %s (%d) je mimo rozsah 1..%d" % (jmeno, sl, pocet))

    ocekavane = {vystup.S_KREDITOR: "Kreditor", vystup.S_KOD: "Kód kategorie",
                 vystup.S_KATEG: "Kategorie", vystup.S_SKUP: "Skupina",
                 vystup.S_ICT: "ICT relevance", vystup.S_POPIS: "Co dodává (LLM)",
                 vystup.S_PRISTUP: "Přístup k datům/systémům",
                 vystup.S_NAHRAD: "Nahraditelnost", vystup.S_KRIT: "Kritičnost",
                 vystup.S_ISO: "Režim dle ISO 27001",
                 vystup.S_VLASTNIK: "Vlastník vztahu",
                 vystup.S_POSOUZENI: "Datum posouzení",
                 vystup.S_PREZKOUM: "Datum příštího přezkoumání",
                 vystup.S_POZNAMKA: "Poznámka"}
    for sl, nazev in ocekavane.items():
        skutecny = vystup.SLOUPCE[sl - 1][0]
        if skutecny != nazev:
            chyby.append("sloupec %d je %r, čekán %r" % (sl, skutecny, nazev))

    # 4) hodnoty v rozbalovacích seznamech musí odpovídat vzorcům
    vzorec = vystup._vzorec_kriticnost(2)
    # jen hodnoty, se kterymi se doopravdy POROVNAVA (za '='), ne navratove texty
    for hodnota in re.findall(r'=\s*"([^"]*)"', vzorec):
        if hodnota in ("", "ano"):
            continue
        if hodnota not in vystup.PRISTUP + vystup.NAHRADITELNOST:
            chyby.append("vzorec porovnává s %r, což není v žádné nabídce" % hodnota)
    # Volby, ktere kriticnost zvedaji, musi byt ve vzorci jmenovite. ZADNY,
    # FYZICKY a BEZNE ve vzorci nejsou zamerne - jsou to "nic se nedeje" vetve,
    # ktere propadnou do BEZNY.
    for hodnota in vystup.PRISTUP + vystup.NAHRADITELNOST:
        if hodnota not in (ZADNY, FYZICKY, BEZNE) and '"%s"' % hodnota not in vzorec:
            chyby.append("nabídka obsahuje %r, ale vzorec s tím nepočítá" % hodnota)
    # Excel bere seznam v datove validaci jen do 255 znaku vcetne uvozovek
    for jmeno, hodnoty in (("PRISTUP", vystup.PRISTUP),
                           ("NAHRADITELNOST", vystup.NAHRADITELNOST)):
        delka = len(",".join(hodnoty)) + 2
        if delka > 255:
            chyby.append("nabídka %s má %d znaků, Excel bere 255" % (jmeno, delka))

    # 5) řídicí znaky z webů a rejstříků nesmí projít do buňky - XLSX je
    #    nedovoluje a openpyxl na nich padal až při ukládání, po celém běhu
    spinave = "ACME\x01 s.r.o.\x07\x1f test\x0b\x9d"
    ciste = vystup._bunka(spinave)
    if vystup._RIDICI.search(ciste):
        chyby.append("_bunka nechala řídicí znak v %r" % ciste)
    if "ACME" not in ciste or "test" not in ciste:
        chyby.append("_bunka zahodila i čitelný text: %r" % ciste)
    dlouhy = vystup._bunka("x" * 40000)
    if len(dlouhy) > 32767:
        chyby.append("_bunka nezkrátila text na limit Excelu (%d znaků)" % len(dlouhy))
    if vystup._bunka(7) != 7 or vystup._bunka(None) is not None:
        chyby.append("_bunka mrší nestringové hodnoty")

    # 6) ICT relevance musí mít každá kategorie z číselníku
    bez_ict = [k for k in KATEGORIE if k != "XXX-00" and not ict_relevance(k)]
    if bez_ict:
        chyby.append("kategorie bez ICT příznaku: %s" % ", ".join(sorted(bez_ict)))

    if chyby:
        print("CHYBY (%d):" % len(chyby))
        for c in chyby:
            print("  ✗ %s" % c)
        return 1
    print("OK: %d případů kritičnosti, syntaxe vzorců, indexy sloupců, "
          "číselník (%d kategorií), čtení odpovědí z chatu, značení OSVČ, "
          "čištění buněk, sazba sešitu" % (len(PRIPADY), len(KATEGORIE)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
