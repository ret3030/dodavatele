"""
Vystupy nastroje: davky pro LLM chat, nacteni jeho odpovedi a Excel se vzorci.

Excel je zamerne "zivy" - kategorie, skupina a ICT relevance nejsou zapsane
hodnoty, ale VLOOKUP do listu Ciselnik. Kdyz nekdo opravi kod kategorie,
prepocita se nazev, skupina, ICT priznak, kriticnost i rezim dle ISO 27001.

Sesit je zaroven pracovni nastroj pro klienta, ne jen export: kazdy list zacina
kartou (nadpis + k cemu list je a jak se pouziva) a prvni list je kosilka.
Proto data na listech nezacinaji na radku 1 - viz _karta().
"""

import csv
import json
import os
import re

from taxonomie import KATEGORIE, VYCHOZI_KOD, ict_relevance

DAVKA = 250            # firem na jeden soubor pro chat
_KOD_RE = re.compile(r"\b[A-Z]{2,4}-\d{2}\b")

# Hodnoty rucne vyplnovanych sloupcu. Musi presne sedet na vzorce nize.
#
# Pristup je zamerne rozdeleny na dve urovne "virtualniho" pristupu. Samotne
# "virtualne" nic nerika - dodavatel, ktery od nas cte data pres API, je neco
# jineho nez dodavatel, ktery nam spravuje servery. Prvni je rizikovy jen
# tehdy, kdyz jde o ICT dodavku, druhy vzdycky.
PRISTUP = [
    "Žádný",                        # nevidí naše data ani nechodí do prostor
    "Fyzický (vstup do prostor)",   # úklid, servis, ostraha, stěhování
    "Virtuální (data a rozhraní)",  # zpracovává naše data, API, cloud, SaaS
    "Virtuální (správa systémů)",   # administrátorský/privilegovaný přístup
    "Fyzický i virtuální",          # obojí zároveň
]
NAHRADITELNOST = [
    "Běžně nahraditelný",           # na trhu je víc alternativ, přechod v týdnech
    "Obtížně nahraditelný",         # náhrada existuje, ale znamená migraci a měsíce
    "Kritická závislost",           # prakticky nenahraditelný, výpadek zastaví provoz
]

SLOUPCE = [
    ("Kreditor", 12),
    ("Název ze vstupu", 34), ("Ověřený název", 34), ("Jistota identity", 15),
    ("IČO", 11), ("DIČ", 14), ("Země", 6), ("Město", 18), ("Web", 26),
    ("NACE", 14), ("Objem ze vstupu", 14),
    ("Kód kategorie", 13), ("Kategorie", 34), ("Skupina", 30), ("ICT relevance", 13),
    ("Co dodává (LLM)", 44), ("Jistota zařazení", 14),
    ("Přístup k datům/systémům", 26), ("Nahraditelnost", 22),
    ("Kritičnost", 16), ("Režim dle ISO 27001", 52),
    ("Poznámky nástroje", 40),
    # revizni blok - vyplnuje klient u kritickych a vyznamnych dodavatelu
    ("Vlastník vztahu", 22), ("Datum posouzení", 16),
    ("Datum příštího přezkoumání", 22), ("Poznámka", 44),
]
# indexy (1-based) klicovych sloupcu - at se vzorce nerozbiji pri zmene poradi
S_KREDITOR, S_POPIS = 1, 16
S_KOD, S_KATEG, S_SKUP, S_ICT = 12, 13, 14, 15
S_PRISTUP, S_NAHRAD = 18, 19
S_KRIT, S_ISO = 20, 21
S_VLASTNIK, S_POSOUZENI, S_PREZKOUM, S_POZNAMKA = 23, 24, 25, 26
RUCNI = (S_PRISTUP, S_NAHRAD, S_VLASTNIK, S_POSOUZENI, S_PREZKOUM, S_POZNAMKA)

# Popis kazdeho sloupce listu Dodavatele a jeho vazba na ISO/IEC 27001, pro
# list Metodika. Poradi a delka musi sedet na SLOUPCE - kdyz nekdo prida
# sloupec a zapomene na tohle, zapis_excel to shodi hned pri generovani.
METODIKA = [
    ("Kreditor", "Identifikátor dodavatele z účetnictví (číslo kreditora). "
     "Když ho vstup neobsahoval, je tu pořadové číslo řádku – to je zároveň "
     "ID, pod kterým firma vystupuje v dávkách pro chat a na listu Podklady.",
     "A.5.19 – evidence dodavatelů musí být jednoznačně adresovatelná"),
    ("Název ze vstupu", "Název dodavatele přesně tak, jak přišel ze vstupního "
     "souboru (z účetnictví).", ""),
    ("Ověřený název", "Úřední název z rejstříku (ARES / RPO SR) nebo VIES, "
     "pokud se liší od vstupu.", "A.5.19 – dodavatele je třeba jednoznačně "
     "identifikovat"),
    ("Jistota identity", "Jak spolehlivě je identita ověřená: rejstřík (podle "
     "IČO nebo přesné shody jména) > VIES (DIČ) > podle názvu > neověřeno.",
     "A.5.19"),
    ("IČO", "Identifikátor dohledaný v ARES / RPO SR.", ""),
    ("DIČ", "Identifikátor použitý k ověření přes VIES u zahraničních "
     "dodavatelů v EU.", ""),
    ("Země", "Sídlo dodavatele.", ""),
    ("Město", "Sídlo dodavatele.", ""),
    ("Web", "Ověřený web dodavatele – ne jen odhad z názvu.", ""),
    ("NACE", "Úředně zapsaný obor podnikání. Jen vodítko pro kategorizaci, "
     "ne finální zařazení – bývá zastaralý nebo obecný.", ""),
    ("Objem ze vstupu", "Obrat nebo částka, pokud byla ve vstupním souboru. "
     "Jen informace pro kontext – do kritičnosti záměrně nevstupuje.", ""),
    ("Kód kategorie", "Kód z Číselníku, který LLM přiřadil podle toho, čím "
     "se dodavatel reálně zabývá.", "A.5.19 – podklad pro posouzení druhu "
     "dodávky"),
    ("Kategorie", "Název kategorie – VLOOKUP do listu Číselník podle kódu.", ""),
    ("Skupina", "Nadřazená skupina kategorie – VLOOKUP do listu Číselník.", ""),
    ("ICT relevance", "ano / hraniční / ne – jestli dodavatel typicky "
     "zpracovává informace firmy nebo se připojuje do jejích systémů. "
     "VLOOKUP do Číselníku.", "A.5.21 – řízení bezpečnosti v ICT "
     "dodavatelském řetězci"),
    ("Co dodává (LLM)", "Stručný popis od LLM, aby šlo zkontrolovat, že kód "
     "kategorie sedí.", ""),
    ("Jistota zařazení", "vysoká / střední / nízká – jak jistý si LLM byl "
     "při zařazení do kategorie.", ""),
    ("Přístup k datům/systémům", "RUČNĚ. Žádný = nevidí naše data ani nechodí "
     "do prostor. Fyzický = vstup do prostor (úklid, servis, ostraha). "
     "Virtuální (data a rozhraní) = zpracovává nebo přenáší naše data, včetně "
     "přístupu přes API, cloud nebo SaaS – sám o sobě to kritického dodavatele "
     "nedělá. Virtuální (správa systémů) = administrátorský či privilegovaný "
     "přístup do našich systémů, vzdálená správa. Fyzický i virtuální = obojí "
     "zároveň. Rozhoduje skutečný stav, ne to, co je ve smlouvě.",
     "A.5.19, A.5.20; A.8.2 (privilegovaná přístupová práva)"),
    ("Nahraditelnost", "RUČNĚ. Běžně nahraditelný = na trhu je víc alternativ, "
     "přechod je otázka týdnů. Obtížně nahraditelný = náhrada existuje, ale "
     "znamená migraci dat, integrace nebo měsíce práce. Kritická závislost = "
     "prakticky nenahraditelný, jeho výpadek zastaví provoz.",
     "A.5.21; A.5.29/A.5.30 (kontinuita provozu)"),
    ("Kritičnost", "KRITICKÝ / VÝZNAMNÝ / BĚŽNÝ – vzorec z ICT relevance, "
     "Přístupu a Nahraditelnosti. Rozhoduje přístup a závislost, ne cena ani "
     "objem fakturace.", "shrnutí A.5.19–A.5.22"),
    ("Režim dle ISO 27001", "Doporučená opatření podle Kritičnosti a ICT "
     "relevance – prověření, bezpečnostní požadavky ve smlouvě, DPA, právo "
     "auditu, monitoring.", "A.5.19–A.5.22"),
    ("Poznámky nástroje", "Technické poznámky z dohledávání (např. web "
     "nenalezen) – informace o procesu, ne o dodavateli samotném.", ""),
    ("Vlastník vztahu", "RUČNĚ: jméno člověka u nás, který za vztah "
     "s dodavatelem odpovídá. U kritického dodavatele nesmí zůstat prázdné – "
     "bez konkrétní osoby se opatření nikdy nevymáhají.",
     "A.5.19; A.5.2 (role a odpovědnosti v bezpečnosti informací)"),
    ("Datum posouzení", "RUČNĚ: kdy bylo naposledy posouzeno riziko tohohle "
     "dodavatele a ověřeno, že sloupce vlevo platí.", "A.5.19, A.5.22 – "
     "monitorování a přezkoumání dodavatelských služeb"),
    ("Datum příštího přezkoumání", "RUČNĚ: kdy se má posouzení zopakovat. "
     "Doporučení: kritický dodavatel po roce, významný po dvou letech, "
     "běžný při změně vztahu. Prošlé datum list zvýrazní červeně.",
     "A.5.22 – přezkoumání v plánovaných intervalech"),
    ("Poznámka", "RUČNĚ: cokoli k tomuhle dodavateli – co bylo prověřeno, "
     "kde je smlouva a DPA, otevřené body, dohodnutá opatření.", "A.5.20"),
]
assert len(METODIKA) == len(SLOUPCE), "METODIKA musí popisovat každý sloupec SLOUPCE"


def _pismeno(i):
    """1 -> A, 27 -> AA"""
    s = ""
    while i:
        i, z = divmod(i - 1, 26)
        s = chr(65 + z) + s
    return s


# XLSX v bunce nedovoluje ridici znaky. Lezou sem z rozbitych webu, z rejstriku
# i ze vstupniho CSV a openpyxl na nich spadne az pri wb.save(), tedy po celem
# behu - proto se sem sahne u kazde hodnoty, ne jen u tech z internetu.
_RIDICI = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f"
                     "\ud800-\udfff\ufdd0-\ufdef\ufffe\uffff]")
LIMIT_BUNKY = 32000   # tvrdy strop Excelu je 32767 znaku


def _bunka(h):
    if not isinstance(h, str):
        return h
    h = _RIDICI.sub(" ", h)
    return h[:LIMIT_BUNKY - 1] + "…" if len(h) > LIMIT_BUNKY else h


def _radek(ws, hodnoty):
    """append() s ocistenim - jediny vstup dat do sesitu. Vraci cislo radku."""
    ws.append([_bunka(h) for h in hodnoty])
    return ws.max_row


# ---------------------------------------------------------------------------
# Davky pro LLM chat
# ---------------------------------------------------------------------------

def _ciselnik_radky():
    radky, predchozi = [], None
    for kod, skupina, nazev in sorted(
            (k, v[0], v[1]) for k, v in KATEGORIE.items() if k != VYCHOZI_KOD):
        if skupina != predchozi:
            radky.append("  [%s]" % skupina)
            predchozi = skupina
        radky.append("    %s = %s" % (kod, nazev))
    return radky


ZADANI = """\
Jsi asistent pro audit dodavatelů podle ISO/IEC 27001 (příloha A, řízení
dodavatelských vztahů). Ke každé firmě níže urči:

  1. KÓD KATEGORIE z číselníku - podle toho, ČÍM SE FIRMA REÁLNĚ ZABÝVÁ,
     ne podle formálního zápisu v rejstříku. Zapsané obory jsou často
     zastaralé nebo obecné ("výroba, obchod a služby") - ber je jen jako
     vodítko. Když firmu znáš, použij i vlastní znalosti.
  2. CO DODÁVÁ - stručně vlastními slovy (pár slov až věta), ať jde
     zkontrolovat, že kód sedí.
  3. JISTOTU - vysoká / střední / nízká. Napiš "nízká" všude, kde odhaduješ
     z názvu nebo si nejsi jistý/á. Raději přiznaná nejistota než tichý omyl -
     na tenhle výstup navazuje audit.

Když o firmě nic nevíš, napiš kód XXX-00 a jistotu "nízká".

U každé firmy je v odrážkách kontext, který se podařilo dohledat ve veřejných
rejstřících a na webu. Pozor na řádek "identita" - hodnota "neověřeno" znamená,
že firmu nemáme potvrzenou z rejstříku a i název může být nepřesný.
"""

FORMAT = """\
Odpověz POUZE v CSV, jeden řádek na firmu, oddělovač ';', v tomto pořadí.
Bez hlavičky, bez uvozovek, bez textu okolo:

ID;Název firmy;Kód kategorie;Co dodává;Jistota

ID i název opiš přesně tak, jak jsou v seznamu - podle nich se odpověď páruje
zpět a nesoulad se pozná. Zpracuj všech %d firem. Když nestihneš všechny,
radši méně řádků, ale úplných - nikdy nevynechávej ID uprostřed.
"""


def _blok_firmy(cislo, f):
    hlava = "%d | %s" % (cislo, f.nazev or f.vstup_nazev)
    misto = ", ".join(x for x in (f.mesto, f.zeme) if x)
    if misto:
        hlava += " (%s)" % misto
    radky = [hlava, "   identita: %s" % f.identita]
    for stitek, text in f.podklady(strucne=True):
        radky.append("   %s: %s" % (stitek, text))
    return "\n".join(radky)


def zapis_davky(firmy, adresar):
    """Rozdeli firmy do souboru po DAVKA kusech. Vraci seznam cest."""
    os.makedirs(adresar, exist_ok=True)
    for star in os.listdir(adresar):
        if re.match(r"^davka_\d+\.txt$", star):
            os.remove(os.path.join(adresar, star))

    ciselnik = "\n".join(_ciselnik_radky())
    davky = [firmy[i:i + DAVKA] for i in range(0, len(firmy), DAVKA)] or [[]]
    cesty = []
    for i, davka in enumerate(davky, 1):
        cislo_od = (i - 1) * DAVKA + 1
        obsah = "\n".join([
            ZADANI,
            "ČÍSELNÍK KATEGORIÍ:",
            ciselnik,
            "",
            "FIRMY (%d, dávka %d z %d):" % (len(davka), i, len(davky)),
            "",
            "\n\n".join(_blok_firmy(cislo_od + j, f) for j, f in enumerate(davka)),
            "",
            FORMAT % len(davka),
        ])
        cesta = os.path.join(adresar, "davka_%02d.txt" % i)
        with open(cesta, "w", encoding="utf-8") as fh:
            fh.write(obsah)
        cesty.append(cesta)
    return cesty


# ---------------------------------------------------------------------------
# Nacteni odpovedi z chatu
# ---------------------------------------------------------------------------

def nacti_odpovedi(adresar, firmy=None):
    """
    Precte vsechny CSV/TXT v adresari a vrati (odpovedi, nesoulady), kde
    odpovedi = {id: {"kod","popis","jistota"}}.

    Radky bez rozpoznatelneho ID nebo kodu se preskoci - chat casto pripise
    uvodni vetu nebo hlavicku. Kdyz je zadany `firmy`, kontroluje se navic
    nazev: kdyby chat posunul cislovani, dostala by firma cizi kategorii
    a v auditu by se to poznalo hodne pozde. Takove radky se zahodi a vrati
    v `nesoulady`.
    """
    out, nesoulady = {}, []
    if not os.path.isdir(adresar):
        return out, nesoulady
    for jmeno in sorted(os.listdir(adresar)):
        if not jmeno.lower().endswith((".csv", ".txt")):
            continue
        with open(os.path.join(adresar, jmeno), encoding="utf-8-sig", newline="") as fh:
            vzorek = fh.read(4096)
            fh.seek(0)
            odd = ";" if vzorek.count(";") >= vzorek.count(",") else ","
            for radek in csv.reader(fh, delimiter=odd):
                if len(radek) < 2:
                    continue
                ident = re.match(r"^\s*(\d+)\s*$", (radek[0] or "").strip())
                if not ident:
                    continue
                # kod kategorie muze byt ve 2. nebo 3. sloupci podle toho,
                # jestli chat vratil i nazev firmy
                kod, i_kod = None, None
                for i in (1, 2):
                    if i < len(radek):
                        m = _KOD_RE.search((radek[i] or "").upper())
                        if m:
                            kod, i_kod = m.group(0), i
                            break
                if not kod:
                    continue

                cislo = int(ident.group(1))
                nazev_z_odpovedi = radek[1].strip() if i_kod == 2 else ""
                if firmy and nazev_z_odpovedi and not _nazev_sedi(
                        firmy, cislo, nazev_z_odpovedi):
                    ocekavany = (firmy[cislo - 1].nazev or firmy[cislo - 1].vstup_nazev
                                 if 1 <= cislo <= len(firmy) else "?")
                    nesoulady.append((cislo, nazev_z_odpovedi, ocekavany))
                    continue

                out[cislo] = {
                    "kod": kod,
                    "popis": (radek[i_kod + 1].strip() if len(radek) > i_kod + 1 else ""),
                    "jistota": (radek[i_kod + 2].strip().lower()
                                if len(radek) > i_kod + 2 else ""),
                }
    return out, nesoulady


def _nazev_sedi(firmy, cislo, nazev):
    """Odpovida nazev z odpovedi firme na tomhle poradovem cisle?"""
    if not 1 <= cislo <= len(firmy):
        return False
    f = firmy[cislo - 1]
    from zdroje import skore_shody
    return max(skore_shody(nazev, f.vstup_nazev),
               skore_shody(nazev, f.nazev or f.vstup_nazev)) >= 0.75


# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

BARVA_TMAVA = "1F3864"      # zahlavi tabulek
BARVA_KARTA = "DDEBF7"      # karta nad tabulkou
BARVA_RUCNI = "FFF2CC"      # bunky, ktere vyplnuje clovek
BARVA_RUCNI_HLAVA = "BF8F00"
BARVA_VYSTRAHA = "FFC7CE"   # chybejici udaj u kritickeho dodavatele
BARVA_KRIT = "C00000"

VYCHOZI_SIRKA = 8.43        # sirka sloupce, kteremu ji nikdo nenastavil


def _sirka(ws, i):
    pis = _pismeno(i)
    if pis in ws.column_dimensions:
        return ws.column_dimensions[pis].width or VYCHOZI_SIRKA
    return VYCHOZI_SIRKA


def _karta(ws, nadpis, radky):
    """
    Vysvetlivka nad tabulkou: nadpis + par radku o tom, co list je a jak se
    pouziva. Vraci cislo radku, na ktery ma prijit zahlavi tabulky.

    Kvuli karte data na listech nezacinaji na radku 1 - vsechny vzorce a
    rozsahy proto pocitaji s navratovou hodnotou, nikde nejsou natvrdo.

    Text se slouci pres tolik sloupcu, aby se nejdelsi radek vesel - ve
    sloucene bunce se text nepretece do sousedni, takze by se jinak orizl.
    Sirky sloupcu proto musi byt nastavene uz pred volanim.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    vypln = PatternFill("solid", fgColor=BARVA_KARTA)
    potreba = max(len(t) for t in [nadpis] + list(radky)) + 2
    sloupcu, sirka = 0, 0.0
    while sirka < potreba and sloupcu < 40:
        sloupcu += 1
        sirka += _sirka(ws, sloupcu)
    posledni_sl = _pismeno(sloupcu)

    def _pruh(r, text, font, vyska):
        ws.cell(r, 1, text).font = font
        ws.cell(r, 1).alignment = Alignment(vertical="center")
        for c in range(1, sloupcu + 1):
            ws.cell(r, c).fill = vypln
        ws.merge_cells("A%d:%s%d" % (r, posledni_sl, r))
        ws.row_dimensions[r].height = vyska

    _pruh(1, nadpis, Font(bold=True, size=13, color=BARVA_TMAVA), 24)
    r = 1
    for text in radky:
        r += 1
        _pruh(r, text, Font(size=10, color=BARVA_TMAVA), 15)
    _pruh(r + 1, "", Font(size=10), 6)
    return r + 3          # jeden prazdny radek mezi kartou a zahlavim


def _vzorec_kriticnost(r):
    """
    Kriticnost dodavatele. Rozhoduji dve veci: jaky ma dodavatel pristup
    (ISO 27001 A.5.19-A.5.22) a jak snadno jde nahradit. Objem fakturace ani
    typ smlouvy do vypoctu nevstupuji - levny dodavatel se vzdalenou spravou
    serveru je rizikovejsi nez drahy dodavatel kancelarskych potreb.

    Pristup pres API nebo do cloudu ("data a rozhrani") sam o sobe kritickeho
    dodavatele nedela - zvedne se na KRITICKY az u ICT dodavky. Spravu systemu
    (privilegovany pristup) bereme jako kritickou vzdycky.
    """
    kod, ict = "$%s%d" % (_pismeno(S_KOD), r), "$%s%d" % (_pismeno(S_ICT), r)
    pri, nah = "$%s%d" % (_pismeno(S_PRISTUP), r), "$%s%d" % (_pismeno(S_NAHRAD), r)
    return (
        '=IF({kod}="","nezařazeno",'
        'IF({pri}="","⟵ doplňte přístup",'
        'IF(OR({pri}="Virtuální (správa systémů)",'
        '{nah}="Kritická závislost",'
        'AND({ict}="ano",OR({pri}="Virtuální (data a rozhraní)",'
        '{pri}="Fyzický i virtuální"))),"KRITICKÝ",'
        'IF(OR({ict}="ano",{pri}="Virtuální (data a rozhraní)",'
        '{pri}="Fyzický i virtuální",{nah}="Obtížně nahraditelný"),'
        '"VÝZNAMNÝ","BĚŽNÝ"))))'
    ).format(kod=kod, ict=ict, pri=pri, nah=nah)


def _vzorec_iso(r):
    krit, ict = "$%s%d" % (_pismeno(S_KRIT), r), "$%s%d" % (_pismeno(S_ICT), r)
    return (
        '=IF({krit}="KRITICKÝ",IF({ict}="ano",'
        '"Prověření dodavatele, bezpečnostní požadavky ve smlouvě vč. DPA, '
        'právo auditu, průběžný monitoring",'
        '"Prověření dodavatele, smluvní záruky, průběžný monitoring"),'
        'IF({krit}="VÝZNAMNÝ","Zjednodušené prověření, bezpečnostní doložka ve smlouvě",'
        'IF({krit}="BĚŽNÝ","Vedení v evidenci, bez dalších požadavků","")))'
    ).format(krit=krit, ict=ict)


def _vlookup(r, sloupec):
    # nazev listu v apostrofech - s diakritikou to nekterym tabulkovym
    # procesorum bez uvozeni vadi
    return "=IFERROR(VLOOKUP($%s%d,'Číselník'!$A:$D,%d,FALSE),\"\")" % (
        _pismeno(S_KOD), r, sloupec)


def _list_uvod(wb, firmy):
    """Kosilka: co sesit je, jak se s nim pracuje a co znamenaji barvy."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.active
    ws.title = "Úvod"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 106
    stav = [0]

    def _pis(a="", b="", tucne=False, vypln=None, vyska=None):
        stav[0] += 1
        r = stav[0]
        bunka_a = ws.cell(r, 1, a)
        bunka_a.font = Font(bold=tucne, size=11,
                            color=BARVA_TMAVA if tucne else "000000")
        bunka_a.alignment = Alignment(vertical="center", wrap_text=True)
        if vypln:
            bunka_a.fill = PatternFill("solid", fgColor=vypln)
        if b:
            bunka_b = ws.cell(r, 2, b)
            bunka_b.alignment = Alignment(vertical="center", wrap_text=True)
        elif a:
            ws.merge_cells("A%d:B%d" % (r, r))
        if vyska:
            ws.row_dimensions[r].height = vyska
        return r

    def _nadpis(text):
        _pis()
        r = _pis(text, tucne=True, vyska=22)
        ws.cell(r, 1).font = Font(bold=True, size=12, color=BARVA_TMAVA)

    r = _pis("Evidence dodavatelů – audit dodavatelských vztahů dle ISO/IEC 27001",
               vyska=30)
    ws.cell(r, 1).font = Font(bold=True, size=16, color=BARVA_TMAVA)
    _pis("Podklad pro řízení dodavatelských vztahů (ISO/IEC 27001, příloha A, "
           "A.5.19–A.5.22)")
    _pis("Dodavatelů v sešitu: %d    ·    vygenerováno %s"
           % (len(firmy), _dnes()))

    _nadpis("K čemu sešit je")
    for text in (
        "Seznam kreditorů z účetnictví převedený na evidenci dodavatelů: kdo "
        "dodavatel doopravdy je, čím se zabývá,",
        "jestli má vazbu na ICT a jak je pro vaši firmu kritický. Identitu "
        "dohledal nástroj ve veřejných rejstřících,",
        "zařazení do kategorie navrhl jazykový model. Vztah k vaší firmě "
        "posoudíte vy – nástroj ho odjinud zjistit neumí.",
        "",
        "Sešit je živý dokument, ne export. Sloupce se vzorci se přepočítají, "
        "jakmile změníte vstup, ze kterého počítají.",
    ):
        _pis(text)

    _nadpis("Jak s ním pracovat")
    for text in (
        "1.  Projděte list Dodavatelé a zkontrolujte zařazení – hlavně řádky, "
        "kde je Jistota identity „neověřeno“ nebo Jistota zařazení „nízká“.",
        "2.  Kde kategorie nesedí, přepište Kód kategorie podle listu Číselník. "
        "Kategorie, Skupina i ICT relevance se dopočítají samy.",
        "3.  U každého dodavatele vyplňte dva žluté sloupce: Přístup "
        "k datům/systémům a Nahraditelnost. Z nich vyjde Kritičnost.",
        "4.  U dodavatelů, kteří vyjdou jako KRITICKÝ (a podle uvážení "
        "i VÝZNAMNÝ), doplňte revizní blok na konci řádku:",
        "     Vlastník vztahu, Datum posouzení, Datum příštího přezkoumání "
        "a Poznámku. Chybějící údaje se zvýrazní červeně.",
        "5.  Sledujte list Souhrn – ukazuje, kolik toho zbývá vyplnit a kolik "
        "dodavatelů čeká na přezkoumání.",
        "",
        "Význam jednotlivých voleb je popsaný na listu Metodika, včetně toho, "
        "ke kterému řízení ISO/IEC 27001 se sloupec váže.",
    ):
        _pis(text)

    _nadpis("Co je na kterém listu")
    for nazev, popis in LISTY:
        r = _pis(nazev, popis)
        ws.cell(r, 1).font = Font(bold=True, size=11)

    _nadpis("Barvy a značky")
    for barva, popis, vysvetleni in (
        (BARVA_TMAVA, "tmavě modré záhlaví", "sloupec vyplní nástroj nebo vzorec – "
         "přepisovat ho nemá smysl, přepíše se zpátky"),
        (BARVA_RUCNI_HLAVA, "žluté záhlaví a buňky", "sem patří vaše rozhodnutí; "
         "buňky s rozbalovacím seznamem berou jen hodnoty z nabídky"),
        (BARVA_VYSTRAHA, "červené podbarvení", "u kritického dodavatele chybí "
         "vlastník nebo datum posouzení, nebo je přezkoumání po termínu"),
    ):
        r = _pis(popis, vysvetleni, vypln=barva)
        ws.cell(r, 1).font = Font(bold=True, size=11,
                                  color="FFFFFF" if barva == BARVA_TMAVA else "000000")

    _nadpis("Na co si dát pozor")
    for text in (
        "·  Kritičnost počítá z přístupu a nahraditelnosti, ne z objemu "
        "fakturace. Levný dodavatel se vzdálenou správou serverů je rizikovější",
        "   než drahý dodavatel kancelářských potřeb. Sloupec Objem ze vstupu "
        "je jen kontext.",
        "·  Řádek s Jistotou identity „neověřeno“ stojí jen na názvu ze vstupu "
        "a na webu – ověřte ho, než na něm postavíte rozhodnutí.",
        "·  Zařazení do kategorie navrhl jazykový model. U „nízké“ jistoty "
        "zařazení je omyl pravděpodobný; podklady jsou na listu Podklady.",
        "·  Prázdný Přístup k datům/systémům znamená, že Kritičnost není "
        "spočítaná – v buňce je „⟵ doplňte přístup“, ne BĚŽNÝ.",
    ):
        _pis(text)
    return ws


LISTY = [
    ("Dodavatelé", "hlavní tabulka – jeden řádek na dodavatele; tady se pracuje"),
    ("Číselník", "kategorie ve skupinách + ICT příznak; zdroj pro vzorce "
     "a rozbalovací seznamy"),
    ("Souhrn", "kolik dodavatelů je ICT, kolik kritických, co ještě chybí vyplnit"),
    ("Podklady", "co se o firmě našlo ve veřejných zdrojích – pro zpětnou kontrolu "
     "zařazení"),
    ("Metodika", "co znamená který sloupec a na které řízení ISO/IEC 27001 se váže"),
]


def _dnes():
    from datetime import date
    return date.today().strftime("%d.%m.%Y")


def zapis_excel(firmy, odpovedi, cesta):
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise SystemExit(
            "Zápis XLSX vyžaduje openpyxl:  pip install openpyxl")

    wb = Workbook()
    _list_uvod(wb, firmy)

    # -- list Dodavatelé --------------------------------------------------
    ws = wb.create_sheet("Dodavatelé")
    for i, (_, sirka) in enumerate(SLOUPCE, 1):
        ws.column_dimensions[get_column_letter(i)].width = sirka
    zahlavi = _karta(ws, "Dodavatelé – hlavní tabulka", [
        "Jeden řádek na dodavatele. Bílé sloupce dohledal nástroj, žluté "
        "vyplňujete vy, Kritičnost a Režim dle ISO 27001 se počítají vzorcem.",
        "Vyplňte Přístup k datům/systémům a Nahraditelnost – bez nich se "
        "kritičnost nespočítá. U kritických dodavatelů pak revizní blok "
        "na konci řádku.",
        "Kategorii opravíte přepsáním Kódu kategorie (nabídka je na listu "
        "Číselník); název, skupina i ICT relevance se dopočítají.",
    ])
    prvni = zahlavi + 1

    for i, (nazev, _) in enumerate(SLOUPCE, 1):
        ws.cell(zahlavi, i, nazev)

    for cislo, f in enumerate(firmy, 1):
        odp = odpovedi.get(cislo) or {}
        r = zahlavi + cislo
        _radek(ws, [
            f.vstup_kod or cislo,
            f.vstup_nazev,
            f.nazev if f.nazev != f.vstup_nazev else "",
            f.identita,
            f.ico, f.dic, f.zeme, f.mesto, f.web,
            ", ".join(f.nace[:6]),
            f.vstup_objem,
            odp.get("kod", ""),
            _vlookup(r, 2), _vlookup(r, 3), _vlookup(r, 4),
            odp.get("popis", ""), odp.get("jistota", ""),
            "", "",
            _vzorec_kriticnost(r), _vzorec_iso(r),
            "; ".join(f.poznamky),
            "", "", "", "",
        ])

    posledni = zahlavi + len(firmy)

    # zahlavi
    vypln = PatternFill("solid", fgColor=BARVA_TMAVA)
    for b in ws[zahlavi]:
        b.font = Font(bold=True, color="FFFFFF")
        b.fill = vypln
        b.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[zahlavi].height = 30
    ws.freeze_panes = "C%d" % prvni
    if posledni >= prvni:
        ws.auto_filter.ref = "A%d:%s%d" % (zahlavi, _pismeno(len(SLOUPCE)), posledni)

    # rucne vyplnovane sloupce - jine pozadi, at je videt, co ceka na cloveka
    rucni = PatternFill("solid", fgColor=BARVA_RUCNI)
    for sl in RUCNI:
        ws.cell(zahlavi, sl).fill = PatternFill("solid", fgColor=BARVA_RUCNI_HLAVA)
        for r in range(prvni, posledni + 1):
            ws.cell(r, sl).fill = rucni

    # rozbalovaci seznamy
    for sl, hodnoty in ((S_PRISTUP, PRISTUP), (S_NAHRAD, NAHRADITELNOST)):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(hodnoty),
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add("%s%d:%s%d" % (_pismeno(sl), prvni, _pismeno(sl),
                              max(posledni, prvni)))

    for r in range(prvni, posledni + 1):
        ws.cell(r, S_KRIT).font = Font(bold=True)
        for sl in (S_KATEG, S_POPIS, S_ISO, S_POZNAMKA):
            ws.cell(r, sl).alignment = Alignment(wrap_text=True, vertical="top")
        for sl in (S_POSOUZENI, S_PREZKOUM):
            ws.cell(r, sl).number_format = "DD.MM.YYYY"

    # zvyrazneni - co u kritickeho dodavatele chybi a co je po termínu
    if posledni >= prvni:
        cerveny = PatternFill("solid", fgColor=BARVA_VYSTRAHA)
        krit = "$%s%d" % (_pismeno(S_KRIT), prvni)
        for sl in (S_VLASTNIK, S_POSOUZENI):
            pis = _pismeno(sl)
            ws.conditional_formatting.add(
                "%s%d:%s%d" % (pis, prvni, pis, posledni),
                FormulaRule(formula=['AND(%s="KRITICKÝ",%s%d="")'
                                     % (krit, pis, prvni)], fill=cerveny))
        pis = _pismeno(S_PREZKOUM)
        ws.conditional_formatting.add(
            "%s%d:%s%d" % (pis, prvni, pis, posledni),
            FormulaRule(formula=['AND(%s%d<>"",%s%d<TODAY())'
                                 % (pis, prvni, pis, prvni)], fill=cerveny))
        pis = _pismeno(S_KRIT)
        ws.conditional_formatting.add(
            "%s%d:%s%d" % (pis, prvni, pis, posledni),
            FormulaRule(formula=['%s%d="KRITICKÝ"' % (pis, prvni)],
                        font=Font(bold=True, color=BARVA_KRIT)))

    # -- list Číselník (zdroj pro VLOOKUP i rozbalovací seznamy) -----------
    ws_c = wb.create_sheet("Číselník")
    for i, sirka in enumerate((14, 44, 34, 14), 1):
        ws_c.column_dimensions[get_column_letter(i)].width = sirka
    zahlavi_c = _karta(ws_c, "Číselník kategorií", [
        "Zdroj pro sloupce Kategorie, Skupina a ICT relevance na listu "
        "Dodavatelé – dosazují se sem podle Kódu kategorie.",
        "Když chcete dodavatele přeřadit, najděte tu správný kód a přepište "
        "ho v listu Dodavatelé. Tady nic přepisovat nemusíte.",
        "ICT relevance říká, jestli kategorie typicky znamená přístup "
        "k informacím firmy nebo do jejích systémů (A.5.21).",
    ])
    ws_c.cell(zahlavi_c, 1, "Kód")
    ws_c.cell(zahlavi_c, 2, "Kategorie")
    ws_c.cell(zahlavi_c, 3, "Skupina")
    ws_c.cell(zahlavi_c, 4, "ICT relevance")
    for kod, (skupina, nazev) in sorted(KATEGORIE.items()):
        _radek(ws_c, [kod, nazev, skupina, ict_relevance(kod)])
    for b in ws_c[zahlavi_c]:
        b.font = Font(bold=True, color="FFFFFF")
        b.fill = vypln
    ws_c.freeze_panes = "A%d" % (zahlavi_c + 1)

    # -- list Souhrn ------------------------------------------------------
    ws_s = wb.create_sheet("Souhrn")
    ws_s.column_dimensions["A"].width = 38
    ws_s.column_dimensions["B"].width = 14
    zahlavi_s = _karta(ws_s, "Souhrn", [
        "Kontrolní pohled na celou evidenci. Čísla jsou vzorce – po každé "
        "změně na listu Dodavatelé se přepočítají sama.",
        "Blok „Co zbývá vyplnit“ ukazuje, jak daleko jste s posouzením; "
        "dokud nejsou nuly, evidence není hotová.",
    ])

    def _rozsah(sloupec):
        pis = _pismeno(sloupec)
        return "'Dodavatelé'!$%s$%d:$%s$%d" % (pis, prvni, pis, max(posledni, prvni))

    kr, ic, kod = _rozsah(S_KRIT), _rozsah(S_ICT), _rozsah(S_KOD)
    ident = _rozsah(4)
    vlastnik, posouzeni = _rozsah(S_VLASTNIK), _rozsah(S_POSOUZENI)
    prezkoum = _rozsah(S_PREZKOUM)

    radky_souhrnu = [
        ("Přehled", ""),
        ("Dodavatelů celkem", len(firmy)),
        ("Zařazeno LLM", '=COUNTIF(%s,"<>")' % kod),
        ("Nezařazeno (XXX-00)", '=COUNTIF(%s,"XXX-00")' % kod),
        (None, None),
        ("Identita", ""),
        ("rejstřík", '=COUNTIF(%s,"rejstřík")' % ident),
        ("VIES (DIČ)", '=COUNTIF(%s,"VIES (DIČ)")' % ident),
        ("podle názvu", '=COUNTIF(%s,"podle názvu")' % ident),
        ("neověřeno", '=COUNTIF(%s,"neověřeno")' % ident),
        (None, None),
        ("ICT relevance", ""),
        ("ano", '=COUNTIF(%s,"ano")' % ic),
        ("hraniční", '=COUNTIF(%s,"hraniční")' % ic),
        ("ne", '=COUNTIF(%s,"ne")' % ic),
        (None, None),
        ("Kritičnost", ""),
        ("KRITICKÝ", '=COUNTIF(%s,"KRITICKÝ")' % kr),
        ("VÝZNAMNÝ", '=COUNTIF(%s,"VÝZNAMNÝ")' % kr),
        ("BĚŽNÝ", '=COUNTIF(%s,"BĚŽNÝ")' % kr),
        ("Kritických s ICT vazbou", '=COUNTIFS(%s,"KRITICKÝ",%s,"ano")' % (kr, ic)),
        (None, None),
        ("Co zbývá vyplnit", ""),
        ("bez přístupu (kritičnost nespočtena)", '=COUNTIF(%s,"⟵*")' % kr),
        ("kritických bez vlastníka vztahu",
         '=COUNTIFS(%s,"KRITICKÝ",%s,"")' % (kr, vlastnik)),
        ("kritických bez data posouzení",
         '=COUNTIFS(%s,"KRITICKÝ",%s,"")' % (kr, posouzeni)),
        ("přezkoumání po termínu",
         '=COUNTIFS(%s,"<"&TODAY(),%s,"<>")' % (prezkoum, prezkoum)),
    ]
    for popisek, hodnota in radky_souhrnu:
        if popisek is None:
            _radek(ws_s, [])
            continue
        r = _radek(ws_s, [popisek, hodnota])
        if hodnota == "":
            ws_s.cell(r, 1).font = Font(bold=True, size=12, color=BARVA_TMAVA)
    ws_s.sheet_view.showGridLines = False

    # -- list Podklady (proc to LLM zaradilo takhle) ----------------------
    ws_p = wb.create_sheet("Podklady")
    for i, sirka in enumerate((12, 34, 14, 120), 1):
        ws_p.column_dimensions[get_column_letter(i)].width = sirka
    zahlavi_p = _karta(ws_p, "Podklady k zařazení", [
        "Co se o firmě našlo ve veřejných zdrojích – zapsané obory, NACE, "
        "Wikidata, Wikipedie a text z webu.",
        "Slouží ke kontrole: když u dodavatele pochybujete o kategorii, "
        "najděte si ho tu podle Kreditora a přečtěte, z čeho se vycházelo.",
        "Firma bez řádků tady neměla dohledatelné podklady – zařazení stojí "
        "jen na názvu a znalostech modelu.",
    ])
    ws_p.cell(zahlavi_p, 1, "Kreditor")
    ws_p.cell(zahlavi_p, 2, "Název")
    ws_p.cell(zahlavi_p, 3, "Zdroj")
    ws_p.cell(zahlavi_p, 4, "Text")
    for cislo, f in enumerate(firmy, 1):
        for stitek, text in f.podklady():
            _radek(ws_p, [f.vstup_kod or cislo, f.nazev or f.vstup_nazev,
                          stitek, text])
    for b in ws_p[zahlavi_p]:
        b.font = Font(bold=True, color="FFFFFF")
        b.fill = vypln
    ws_p.freeze_panes = "A%d" % (zahlavi_p + 1)

    # -- list Metodika (co který sloupec znamená a vazba na ISO 27001) ----
    ws_m = wb.create_sheet("Metodika")
    for i, sirka in enumerate((28, 78, 40), 1):
        ws_m.column_dimensions[get_column_letter(i)].width = sirka
    zahlavi_m = _karta(ws_m, "Metodika", [
        "Co znamená který sloupec listu Dodavatelé, co do něj patří a na "
        "které řízení ISO/IEC 27001 se váže.",
        "Sem se dívejte, když nevíte, kterou hodnotu z rozbalovacího seznamu "
        "vybrat – u Přístupu a Nahraditelnosti jsou popsané všechny volby.",
        "Prázdná vazba na ISO znamená, že sloupec je jen identifikační nebo "
        "provozní, ne důkaz k žádnému řízení.",
    ])
    ws_m.cell(zahlavi_m, 1, "Sloupec listu Dodavatelé")
    ws_m.cell(zahlavi_m, 2, "Popis")
    ws_m.cell(zahlavi_m, 3, "Vazba na ISO/IEC 27001")
    for nazev, popis, iso in METODIKA:
        _radek(ws_m, [nazev, popis, iso])
    _radek(ws_m, [])
    r = _radek(ws_m, ["Jak se počítá Kritičnost", "", ""])
    ws_m.cell(r, 1).font = Font(bold=True, size=12, color=BARVA_TMAVA)
    for nazev, popis in (
        ("KRITICKÝ", "dodavatel má přístup ke správě systémů, nebo je na něm "
         "kritická závislost, nebo jde o ICT dodávku s přístupem k datům "
         "či do prostor a systémů zároveň"),
        ("VÝZNAMNÝ", "ICT dodávka, nebo přístup k datům a rozhraním, nebo "
         "obtížná nahraditelnost"),
        ("BĚŽNÝ", "vše ostatní – žádný nebo pouze fyzický přístup u běžně "
         "nahraditelného dodavatele mimo ICT"),
        ("⟵ doplňte přístup", "není vyplněný Přístup k datům/systémům; "
         "kritičnost se záměrně nepočítá, aby prázdný řádek nevypadal jako "
         "posouzený"),
    ):
        _radek(ws_m, [nazev, popis, ""])
    for b in ws_m[zahlavi_m]:
        b.font = Font(bold=True, color="FFFFFF")
        b.fill = vypln
    ws_m.freeze_panes = "A%d" % (zahlavi_m + 1)
    for radek in ws_m.iter_rows(min_row=zahlavi_m + 1):
        for bunka in radek:
            bunka.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(cesta)
